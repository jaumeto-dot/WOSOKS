from __future__ import annotations

import base64
import hashlib
import hmac
import json
import os
import re
import secrets
import sqlite3
import tempfile
from datetime import datetime, timedelta, timezone
from html import escape
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, unquote_plus, urlparse
from wsgiref.simple_server import make_server

from openpyxl import load_workbook

try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError:  # Local static use can still run with SQLite only.
    psycopg = None
    dict_row = None


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "wos.sqlite3"
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
ADMIN_PASSWORD_PATH = DATA_DIR / "admin_password.txt"
IMPORTS_DIR = Path(
    os.environ.get("WOS_IMPORTS_DIR")
    or (Path(tempfile.gettempdir()) / "wos-imports" if DATABASE_URL else DATA_DIR / "imports")
)
UPLOAD_XLSX = ROOT / "Hoja Control SHIMA.xlsx"
SOURCE_XLSX = Path("/workspace/scratch/73925ec744f1/upload/Hoja Control SHIMA.xlsx")
OUTLETS = ["SHIMA", "LLUM I SAL", "MEL", "CERCLE", "QUIOSC"]
ROLES = {"viewer", "editor", "admin"}
SESSION_DAYS = 7
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
STATIC_ROOT = ROOT / "web"
DB_READY = False


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def parse_iso(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def hash_password(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 260_000)
    return "pbkdf2_sha256$" + base64.b64encode(salt).decode() + "$" + base64.b64encode(digest).decode()


def verify_password(password: str, stored: str) -> bool:
    try:
        method, salt_text, digest_text = stored.split("$", 2)
        if method != "pbkdf2_sha256":
            return False
        salt = base64.b64decode(salt_text)
        expected = base64.b64decode(digest_text)
        actual = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 260_000)
        return hmac.compare_digest(actual, expected)
    except Exception:
        return False


def bootstrap_admin_password() -> str:
    env_password = os.environ.get("WOS_ADMIN_PASSWORD")
    if env_password:
        return env_password
    DATA_DIR.mkdir(exist_ok=True)
    if ADMIN_PASSWORD_PATH.exists():
        return ADMIN_PASSWORD_PATH.read_text(encoding="utf-8").strip()
    password = secrets.token_urlsafe(12)
    ADMIN_PASSWORD_PATH.write_text(password + "\n", encoding="utf-8")
    return password


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def safe_filename(value: str) -> str:
    filename = Path(value).name.strip() or "import.xlsx"
    filename = re.sub(r"[^A-Za-z0-9._ -]+", "_", filename)
    return filename[:120] or "import.xlsx"


def money(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    text = text.replace("€", "").strip()
    try:
        number = float(str(text).replace(",", "."))
    except ValueError:
        return clean(value)
    if number.is_integer():
        return f"{int(number)}€"
    return f"{number:.2f}".replace(".", ",") + "€"


def norm_label(value: str) -> str:
    return re.sub(r"\s+", " ", clean(value)).strip()


def norm_key(value: str) -> str:
    text = norm_label(value).lower()
    text = (
        text.replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
    )
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_country(value: str) -> str:
    fixes = {
        "espana": "España",
        "esapana": "España",
        "francia": "Francia",
        "japon": "Japón",
    }
    return fixes.get(norm_key(value), norm_label(value))


def normalize_region(value: str) -> str:
    fixes = {
        "galicia": "Galicia",
        "burgundy": "Borgoña",
        "borgona": "Borgoña",
        "rhonevalley": "Rhône Valley",
        "valledelrodano": "Rhône Valley",
    }
    return fixes.get(norm_key(value), norm_label(value))


def normalize_type(value: str, category: str) -> str:
    source = category if norm_key(category) == "sake" else value or category
    fixes = {
        "blanco": "Blancos",
        "blancos": "Blancos",
        "blancosinternacionales": "Blancos Internacionales",
        "botella": "Sake" if norm_key(category) == "sake" else "Botella",
        "dulce": "Dulces",
        "dulces": "Dulces",
        "espumoso": "Espumosos",
        "espumosos": "Espumosos",
        "generoso": "Generosos",
        "generosos": "Generosos",
        "porcopa": "Por copa",
        "rosado": "Rosados",
        "rosados": "Rosados",
        "sake": "Sake",
        "tinto": "Tintos",
        "tintos": "Tintos",
        "tintosinternacionales": "Tintos Internacionales",
    }
    return fixes.get(norm_key(source), norm_label(source))


def canonical(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


FIELD_ALIASES = {
    "name": ["referencia", "vino", "nombre", "wine", "reference"],
    "producer": ["bodegaproductor", "productor", "bodega", "producer"],
    "grapes": ["uvas", "variedades", "variedad", "grape"],
    "vintage": ["añada", "ano", "año", "vintage"],
    "country": ["pais", "país", "country"],
    "region": ["region", "región"],
    "appellation": ["denominacionaoc", "denominacion", "aoc", "do"],
    "subregion": ["subregion", "subregión", "pueblo"],
    "format": ["formato", "format"],
    "price": ["precio", "pvp", "price"],
    "location": ["ubic", "ubicacion", "ubicación", "location"],
    "type": ["tipo", "seccion", "sección", "estilo"],
}


POSTGRES_SCHEMA = """
CREATE TABLE IF NOT EXISTS wines (
  id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
  name TEXT NOT NULL,
  producer TEXT,
  grapes TEXT,
  vintage TEXT,
  country TEXT,
  region TEXT,
  appellation TEXT,
  subregion TEXT,
  format TEXT,
  type TEXT,
  category TEXT,
  search_key TEXT NOT NULL UNIQUE
);

CREATE TABLE IF NOT EXISTS outlets (
  id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
  name TEXT NOT NULL UNIQUE
);

CREATE TABLE IF NOT EXISTS wine_outlets (
  id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
  wine_id BIGINT NOT NULL REFERENCES wines(id) ON DELETE CASCADE,
  outlet_id BIGINT NOT NULL REFERENCES outlets(id) ON DELETE CASCADE,
  location TEXT,
  price TEXT,
  active BOOLEAN NOT NULL DEFAULT TRUE,
  updated_at TIMESTAMPTZ,
  UNIQUE(wine_id, outlet_id)
);

CREATE TABLE IF NOT EXISTS users (
  id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
  username TEXT NOT NULL UNIQUE,
  password_hash TEXT NOT NULL,
  role TEXT NOT NULL DEFAULT 'viewer',
  active BOOLEAN NOT NULL DEFAULT TRUE,
  created_at TIMESTAMPTZ NOT NULL
);

CREATE TABLE IF NOT EXISTS sessions (
  token TEXT PRIMARY KEY,
  user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  expires_at TIMESTAMPTZ NOT NULL,
  created_at TIMESTAMPTZ NOT NULL
);

CREATE TABLE IF NOT EXISTS import_batches (
  id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
  outlet TEXT NOT NULL,
  filename TEXT,
  stored_path TEXT,
  imported_count INTEGER NOT NULL,
  user_id BIGINT REFERENCES users(id),
  created_at TIMESTAMPTZ NOT NULL
);

CREATE INDEX IF NOT EXISTS idx_wine_outlets_active ON wine_outlets(active);
CREATE INDEX IF NOT EXISTS idx_import_batches_created_at ON import_batches(created_at DESC);
"""


def detect_columns(headers: list[str]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    canon_headers = [canonical(h) for h in headers]
    for field, aliases in FIELD_ALIASES.items():
        aliases_canon = [canonical(a) for a in aliases]
        for index, header in enumerate(canon_headers):
            if header in aliases_canon or any(a and a in header for a in aliases_canon):
                mapped[field] = index
                break
    return mapped


def get(row: tuple[Any, ...], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return ""
    return clean(row[index])


def parse_workbook(path: Path, outlet: str) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    records: list[dict[str, str]] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [clean(v) for v in rows[0]]
        mapping = detect_columns(headers)
        if "name" not in mapping:
            continue
        for row in rows[1:]:
            name = get(row, mapping, "name")
            producer = get(row, mapping, "producer")
            grapes = get(row, mapping, "grapes")
            vintage = get(row, mapping, "vintage")
            price = money(row[mapping["price"]]) if "price" in mapping and mapping["price"] < len(row) else ""
            location = get(row, mapping, "location")
            category = normalize_type(ws.title, ws.title)
            wine_type = normalize_type(get(row, mapping, "type") or ws.title, category)

            if not name or name.lower() in {"champagne", "por copa"}:
                continue
            if not any([producer, grapes, vintage, price, location]):
                continue

            records.append(
                {
                    "outlet": outlet,
                    "category": category,
                    "type": wine_type,
                    "name": name,
                    "producer": producer,
                    "grapes": grapes,
                    "vintage": vintage,
                    "country": normalize_country(get(row, mapping, "country")),
                    "region": normalize_region(get(row, mapping, "region")),
                    "appellation": get(row, mapping, "appellation"),
                    "subregion": get(row, mapping, "subregion"),
                    "format": get(row, mapping, "format"),
                    "price": price,
                    "location": location,
                    "active": "1",
                    "updated_at": now_iso(),
                }
            )
    return records


class Database:
    def __init__(self, conn: Any, postgres: bool = False):
        self.conn = conn
        self.postgres = postgres

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        if not self.postgres:
            if exc_type is None:
                self.conn.commit()
            self.conn.close()
        else:
            self.conn.close()

    def execute(self, sql: str, params: tuple[Any, ...] | dict[str, Any] = ()):
        if self.postgres:
            sql = sql.replace("?", "%s")
            if isinstance(params, dict):
                sql = re.sub(r":([A-Za-z_][A-Za-z0-9_]*)", r"%(\1)s", sql)
        return self.conn.execute(sql, params)

    def executescript(self, script: str) -> None:
        if self.postgres:
            self.conn.execute(script)
        else:
            self.conn.executescript(script)


def connect() -> Database:
    if DATABASE_URL:
        if psycopg is None or dict_row is None:
            raise RuntimeError("psycopg no está instalado")
        return Database(psycopg.connect(DATABASE_URL, row_factory=dict_row, autocommit=True), True)
    DATA_DIR.mkdir(exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return Database(conn)


def init_db() -> None:
    with connect() as conn:
        if DATABASE_URL:
            conn.executescript(POSTGRES_SCHEMA)
        else:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS wines (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  name TEXT NOT NULL,
                  producer TEXT,
                  grapes TEXT,
                  vintage TEXT,
                  country TEXT,
                  region TEXT,
                  appellation TEXT,
                  subregion TEXT,
                  format TEXT,
                  type TEXT,
                  category TEXT,
                  search_key TEXT NOT NULL,
                  UNIQUE(search_key)
                );

                CREATE TABLE IF NOT EXISTS outlets (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  name TEXT NOT NULL UNIQUE
                );

                CREATE TABLE IF NOT EXISTS wine_outlets (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  wine_id INTEGER NOT NULL,
                  outlet_id INTEGER NOT NULL,
                  location TEXT,
                  price TEXT,
                  active INTEGER NOT NULL DEFAULT 1,
                  updated_at TEXT,
                  UNIQUE(wine_id, outlet_id),
                  FOREIGN KEY(wine_id) REFERENCES wines(id),
                  FOREIGN KEY(outlet_id) REFERENCES outlets(id)
                );

                CREATE TABLE IF NOT EXISTS users (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  username TEXT NOT NULL UNIQUE,
                  password_hash TEXT NOT NULL,
                  role TEXT NOT NULL DEFAULT 'viewer',
                  active INTEGER NOT NULL DEFAULT 1,
                  created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS sessions (
                  token TEXT PRIMARY KEY,
                  user_id INTEGER NOT NULL,
                  expires_at TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(user_id) REFERENCES users(id)
                );

                CREATE TABLE IF NOT EXISTS import_batches (
                  id INTEGER PRIMARY KEY AUTOINCREMENT,
                  outlet TEXT NOT NULL,
                  filename TEXT,
                  imported_count INTEGER NOT NULL,
                  user_id INTEGER,
                  created_at TEXT NOT NULL,
                  FOREIGN KEY(user_id) REFERENCES users(id)
                );
                """
            )
        for outlet in OUTLETS:
            if DATABASE_URL:
                conn.execute("INSERT INTO outlets(name) VALUES (?) ON CONFLICT (name) DO NOTHING", (outlet,))
            else:
                conn.execute("INSERT OR IGNORE INTO outlets(name) VALUES (?)", (outlet,))
        existing_admin = conn.execute("SELECT id FROM users WHERE username = ?", ("admin",)).fetchone()
        if not existing_admin:
            conn.execute(
                """
                INSERT INTO users(username, password_hash, role, active, created_at)
                VALUES (?, ?, 'admin', TRUE, ?)
                """,
                ("admin", hash_password(bootstrap_admin_password()), now_iso()),
            )


def record_key(record: dict[str, str]) -> str:
    return canonical(
        "|".join(
            [
                record["name"],
                record["producer"],
                record["vintage"],
                record["country"],
                record["region"],
                record["category"],
                record["type"],
                record["format"],
            ]
        )
    )


def import_records(records: list[dict[str, str]], outlet_name: str) -> int:
    with connect() as conn:
        outlet_id = conn.execute("SELECT id FROM outlets WHERE name = ?", (outlet_name,)).fetchone()["id"]
        conn.execute(
            "UPDATE wine_outlets SET active = FALSE, updated_at = ? WHERE outlet_id = ?",
            (now_iso(), outlet_id),
        )
        count = 0
        for record in records:
            key = record_key(record)
            if not key:
                continue
            conn.execute(
                """
                INSERT INTO wines(name, producer, grapes, vintage, country, region, appellation, subregion, format, type, category, search_key)
                VALUES (:name, :producer, :grapes, :vintage, :country, :region, :appellation, :subregion, :format, :type, :category, :search_key)
                ON CONFLICT(search_key) DO UPDATE SET
                  name=excluded.name, producer=excluded.producer, grapes=excluded.grapes, vintage=excluded.vintage,
                  country=excluded.country, region=excluded.region, appellation=excluded.appellation,
                  subregion=excluded.subregion, format=excluded.format, type=excluded.type, category=excluded.category
                """,
                {**record, "search_key": key},
            )
            wine_id = conn.execute("SELECT id FROM wines WHERE search_key = ?", (key,)).fetchone()["id"]
            conn.execute(
                """
                INSERT INTO wine_outlets(wine_id, outlet_id, location, price, active, updated_at)
                VALUES (?, ?, ?, ?, TRUE, ?)
                ON CONFLICT(wine_id, outlet_id) DO UPDATE SET
                  location=excluded.location, price=excluded.price, active=TRUE, updated_at=excluded.updated_at
                """,
                (wine_id, outlet_id, record["location"], record["price"], record["updated_at"]),
            )
            count += 1
        return count


def seed_from_excel() -> None:
    init_db()
    with connect() as conn:
        existing = conn.execute("SELECT COUNT(*) AS n FROM wine_outlets WHERE active = TRUE").fetchone()["n"]
    if SOURCE_XLSX.exists() and not UPLOAD_XLSX.exists():
        UPLOAD_XLSX.write_bytes(SOURCE_XLSX.read_bytes())
    if UPLOAD_XLSX.exists():
        records = parse_workbook(UPLOAD_XLSX, "SHIMA")
        if existing >= len(records):
            return
        import_records(records, "SHIMA")


def ensure_db() -> None:
    global DB_READY
    if DB_READY:
        return
    init_db()
    DB_READY = True


def all_wines() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT w.*, o.name AS outlet, wo.location, wo.price, wo.active, wo.updated_at
            FROM wines w
            JOIN wine_outlets wo ON wo.wine_id = w.id
            JOIN outlets o ON o.id = wo.outlet_id
            WHERE wo.active = TRUE
            ORDER BY lower(w.name)
            """
        ).fetchall()
    return [dict(r) for r in rows]


def grouped_wines() -> list[dict[str, Any]]:
    grouped: dict[int, dict[str, Any]] = {}
    for row in all_wines():
        wine_id = row["id"]
        if wine_id not in grouped:
            grouped[wine_id] = {
                "id": wine_id,
                "name": row["name"],
                "producer": row["producer"],
                "grapes": row["grapes"],
                "vintage": row["vintage"],
                "country": row["country"],
                "region": row["region"],
                "appellation": row["appellation"],
                "subregion": row["subregion"],
                "format": row["format"],
                "type": row["type"],
                "category": row["category"],
                "outlets": [],
            }
        grouped[wine_id]["outlets"].append(
            {
                "outlet": row["outlet"],
                "location": row["location"],
                "price": row["price"],
                "updated_at": row["updated_at"],
            }
        )
    return list(grouped.values())


def cookie_value(environ: dict[str, Any], name: str) -> str:
    cookie_header = environ.get("HTTP_COOKIE", "")
    for part in cookie_header.split(";"):
        if "=" not in part:
            continue
        key, value = part.strip().split("=", 1)
        if key == name:
            return unquote_plus(value)
    return ""


def current_user(environ: dict[str, Any]) -> dict[str, Any] | None:
    token = cookie_value(environ, "wos_session")
    if not token:
        return None
    with connect() as conn:
        row = conn.execute(
            """
            SELECT s.token, s.expires_at, u.id, u.username, u.role, u.active
            FROM sessions s
            JOIN users u ON u.id = s.user_id
            WHERE s.token = ?
            """,
            (token,),
        ).fetchone()
        if not row or not row["active"]:
            return None
        expires_at = row["expires_at"] if isinstance(row["expires_at"], datetime) else parse_iso(row["expires_at"])
        if expires_at <= datetime.now(timezone.utc):
            conn.execute("DELETE FROM sessions WHERE token = ?", (token,))
            return None
        return dict(row)


def create_session(user_id: int) -> str:
    token = secrets.token_urlsafe(32)
    expires_at = datetime.now(timezone.utc) + timedelta(days=SESSION_DAYS)
    with connect() as conn:
        conn.execute(
            "INSERT INTO sessions(token, user_id, expires_at, created_at) VALUES (?, ?, ?, ?)",
            (token, user_id, expires_at.isoformat(timespec="seconds"), now_iso()),
        )
    return token


def clear_session(token: str) -> None:
    if token:
        with connect() as conn:
            conn.execute("DELETE FROM sessions WHERE token = ?", (token,))


def login_page(error: str = "") -> str:
    error_html = f'<p class="error">{escape(error)}</p>' if error else ""
    return f"""<!doctype html>
<html lang="es">
  <head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>WOS · Login</title>
    <style>
      body {{ margin: 0; min-height: 100vh; display: grid; place-items: center; background: #fbfbfa; font-family: Inter, system-ui, sans-serif; color: #050505; }}
      main {{ width: min(420px, calc(100vw - 32px)); background: white; border: 1px solid #dfdfdf; border-radius: 12px; padding: 28px; box-shadow: 0 12px 40px rgba(0,0,0,.08); }}
      h1 {{ margin: 0 0 4px; font-size: 28px; }}
      p {{ margin: 0 0 22px; color: #67717f; }}
      label {{ display: block; margin: 14px 0 6px; font-weight: 700; }}
      input {{ width: 100%; height: 44px; border: 1px solid #dfdfdf; border-radius: 8px; padding: 0 12px; font-size: 16px; box-sizing: border-box; }}
      button {{ width: 100%; height: 46px; margin-top: 20px; border: 0; border-radius: 8px; background: #050505; color: white; font-size: 16px; font-weight: 800; }}
      .error {{ color: #b91c1c; margin: 14px 0 0; }}
    </style>
  </head>
  <body>
    <main>
      <h1>WOS</h1>
      <p>Acceso protegido</p>
      <form method="post" action="/login">
        <label for="username">Usuario</label>
        <input id="username" name="username" autocomplete="username" required />
        <label for="password">Contraseña</label>
        <input id="password" name="password" type="password" autocomplete="current-password" required />
        {error_html}
        <button type="submit">Entrar</button>
      </form>
    </main>
  </body>
</html>"""


def json_response(
    data: Any,
    status: str = "200 OK",
    headers: list[tuple[str, str]] | None = None,
) -> tuple[str, list[tuple[str, str]], bytes]:
    response_headers = [
        ("Content-Type", "application/json; charset=utf-8"),
        ("Cache-Control", "no-store, max-age=0"),
    ]
    if headers:
        response_headers.extend(headers)
    return status, response_headers, json.dumps(data, ensure_ascii=False, default=str).encode()


def database_summary() -> dict[str, Any]:
    with connect() as conn:
        active_rows = conn.execute("SELECT COUNT(*) AS n FROM wine_outlets WHERE active = TRUE").fetchone()["n"]
        wine_rows = conn.execute("SELECT COUNT(*) AS n FROM wines").fetchone()["n"]
        imports = conn.execute(
            """
            SELECT outlet, filename, imported_count, created_at
            FROM import_batches
            ORDER BY created_at DESC
            LIMIT 5
            """
        ).fetchall()
    return {
        "wines": wine_rows,
        "active_locations": active_rows,
        "grouped_wines": len(grouped_wines()),
        "recent_imports": [dict(row) for row in imports],
    }


def parse_json_body(environ: dict[str, Any]) -> dict[str, Any]:
    size = int(environ.get("CONTENT_LENGTH") or "0")
    if size <= 0:
        return {}
    try:
        data = json.loads(environ["wsgi.input"].read(size).decode("utf-8"))
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def html_response(
    html: str,
    status: str = "200 OK",
    headers: list[tuple[str, str]] | None = None,
) -> tuple[str, list[tuple[str, str]], bytes]:
    response_headers = [("Content-Type", "text/html; charset=utf-8")]
    if headers:
        response_headers.extend(headers)
    return status, response_headers, html.encode()


def redirect_response(path: str, headers: list[tuple[str, str]] | None = None) -> tuple[str, list[tuple[str, str]], bytes]:
    response_headers = [("Location", path)]
    if headers:
        response_headers.extend(headers)
    return "302 Found", response_headers, b""


def secure_cookie_suffix(environ: dict[str, Any]) -> str:
    proto = environ.get("HTTP_X_FORWARDED_PROTO", "")
    secure = "; Secure" if proto == "https" else ""
    return f"; Path=/; HttpOnly; SameSite=Lax{secure}"


def require_user(environ: dict[str, Any]) -> tuple[dict[str, Any] | None, tuple[str, list[tuple[str, str]], bytes] | None]:
    user = current_user(environ)
    if user:
        return user, None
    request_path = urlparse(environ.get("PATH_INFO", "/")).path
    if request_path.startswith("/api/"):
        return None, json_response({"error": "No autorizado"}, "401 Unauthorized")
    return None, redirect_response("/login")


def require_admin(environ: dict[str, Any]) -> tuple[dict[str, Any] | None, tuple[str, list[tuple[str, str]], bytes] | None]:
    user, response = require_user(environ)
    if response:
        return None, response
    if user and user["role"] == "admin":
        return user, None
    return None, json_response({"error": "Permiso insuficiente"}, "403 Forbidden")


def require_editor(environ: dict[str, Any]) -> tuple[dict[str, Any] | None, tuple[str, list[tuple[str, str]], bytes] | None]:
    user, response = require_user(environ)
    if response:
        return None, response
    if user and user["role"] in {"admin", "editor"}:
        return user, None
    return None, json_response({"error": "Permiso insuficiente"}, "403 Forbidden")


def list_users() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT id, username, role, active, created_at
            FROM users
            ORDER BY lower(username)
            """
        ).fetchall()
    return [dict(row) for row in rows]


def active_admin_count() -> int:
    with connect() as conn:
        return conn.execute("SELECT COUNT(*) AS n FROM users WHERE role = 'admin' AND active = TRUE").fetchone()["n"]


def create_user(data: dict[str, Any]) -> tuple[str, list[tuple[str, str]], bytes]:
    username = clean(data.get("username", ""))
    password = str(data.get("password", ""))
    role = clean(data.get("role", "viewer"))
    if not re.fullmatch(r"[A-Za-z0-9._-]{3,40}", username):
        return json_response({"error": "Usuario inválido"}, "400 Bad Request")
    if role not in ROLES:
        return json_response({"error": "Rol inválido"}, "400 Bad Request")
    if len(password) < 8:
        return json_response({"error": "Contraseña mínima: 8 caracteres"}, "400 Bad Request")
    try:
        with connect() as conn:
            conn.execute(
                """
                INSERT INTO users(username, password_hash, role, active, created_at)
                VALUES (?, ?, ?, TRUE, ?)
                """,
                (username, hash_password(password), role, now_iso()),
            )
    except Exception:
        return json_response({"error": "Ese usuario ya existe"}, "409 Conflict")
    return json_response({"ok": True, "users": list_users()}, "201 Created")


def update_user(data: dict[str, Any], actor: dict[str, Any]) -> tuple[str, list[tuple[str, str]], bytes]:
    user_id = clean(data.get("id", ""))
    role = clean(data.get("role", ""))
    password = str(data.get("password", ""))
    active = data.get("active", None)
    if not user_id.isdigit():
        return json_response({"error": "Usuario inválido"}, "400 Bad Request")
    with connect() as conn:
        target = conn.execute("SELECT id, username, role, active FROM users WHERE id = ?", (user_id,)).fetchone()
    if not target:
        return json_response({"error": "Usuario no encontrado"}, "404 Not Found")

    will_be_admin = role == "admin" if role else target["role"] == "admin"
    will_be_active = bool(active) if isinstance(active, bool) else bool(target["active"])
    if target["role"] == "admin" and (not will_be_admin or not will_be_active) and active_admin_count() <= 1:
        return json_response({"error": "Debe quedar al menos un admin activo"}, "400 Bad Request")
    if str(target["id"]) == str(actor["id"]) and isinstance(active, bool) and not active:
        return json_response({"error": "No puedes desactivar tu propio usuario"}, "400 Bad Request")

    updates: list[str] = []
    params: list[Any] = []
    if role:
        if role not in ROLES:
            return json_response({"error": "Rol inválido"}, "400 Bad Request")
        updates.append("role = ?")
        params.append(role)
    if isinstance(active, bool):
        updates.append("active = ?")
        params.append(active)
    if password:
        if len(password) < 8:
            return json_response({"error": "Contraseña mínima: 8 caracteres"}, "400 Bad Request")
        updates.append("password_hash = ?")
        params.append(hash_password(password))
    if not updates:
        return json_response({"ok": True, "users": list_users()})
    params.append(user_id)
    with connect() as conn:
        conn.execute(f"UPDATE users SET {', '.join(updates)} WHERE id = ?", tuple(params))
    return json_response({"ok": True, "users": list_users()})


def static_response(path: Path) -> tuple[str, list[tuple[str, str]], bytes]:
    if not path.exists() or not path.is_file():
        return "404 Not Found", [("Content-Type", "text/plain")], b"Not found"
    content_type = "text/html; charset=utf-8"
    if path.suffix == ".css":
        content_type = "text/css; charset=utf-8"
    elif path.suffix == ".js":
        content_type = "application/javascript; charset=utf-8"
    elif path.suffix == ".svg":
        content_type = "image/svg+xml"
    return "200 OK", [("Content-Type", content_type)], path.read_bytes()


def static_file_for(request_path: str) -> Path | None:
    relative = "index.html" if request_path == "/" else request_path.lstrip("/")
    candidate = (STATIC_ROOT / relative).resolve()
    try:
        candidate.relative_to(STATIC_ROOT.resolve())
    except ValueError:
        return None
    return candidate


def app(environ: dict[str, Any], start_response):
    ensure_db()
    request_path = urlparse(environ.get("PATH_INFO", "/")).path
    method = environ.get("REQUEST_METHOD", "GET")

    if request_path == "/login" and method == "GET":
        status, headers, body = html_response(login_page())
    elif request_path == "/login" and method == "POST":
        size = int(environ.get("CONTENT_LENGTH") or "0")
        form = parse_qs(environ["wsgi.input"].read(size).decode())
        username = clean(form.get("username", [""])[0])
        password = form.get("password", [""])[0]
        with connect() as conn:
            row = conn.execute(
                "SELECT id, password_hash, active FROM users WHERE username = ?",
                (username,),
            ).fetchone()
        if row and row["active"] and verify_password(password, row["password_hash"]):
            token = create_session(row["id"])
            cookie = f"wos_session={quote(token)}{secure_cookie_suffix(environ)}; Max-Age={SESSION_DAYS * 24 * 3600}"
            status, headers, body = redirect_response("/", [("Set-Cookie", cookie)])
        else:
            status, headers, body = html_response(login_page("Usuario o contraseña incorrectos"), "401 Unauthorized")
    elif request_path == "/logout":
        clear_session(cookie_value(environ, "wos_session"))
        status, headers, body = redirect_response(
            "/login",
            [("Set-Cookie", f"wos_session={secure_cookie_suffix(environ)}; Max-Age=0")],
        )
    elif request_path == "/api/me":
        user, response = require_user(environ)
        if response:
            status, headers, body = response
        else:
            status, headers, body = json_response({"username": user["username"], "role": user["role"]})
    elif request_path == "/api/wines":
        user, response = require_user(environ)
        if response:
            status, headers, body = response
        else:
            status, headers, body = json_response({"wines": grouped_wines(), "outlets": OUTLETS})
    elif request_path == "/api/diagnostics":
        user, response = require_admin(environ)
        if response:
            status, headers, body = response
        else:
            status, headers, body = json_response(database_summary())
    elif request_path == "/api/users" and method == "GET":
        user, response = require_admin(environ)
        if response:
            status, headers, body = response
        else:
            status, headers, body = json_response({"users": list_users(), "roles": sorted(ROLES)})
    elif request_path == "/api/users" and method == "POST":
        user, response = require_admin(environ)
        if response:
            status, headers, body = response
        else:
            status, headers, body = create_user(parse_json_body(environ))
    elif request_path == "/api/users" and method == "PATCH":
        user, response = require_admin(environ)
        if response:
            status, headers, body = response
        else:
            status, headers, body = update_user(parse_json_body(environ), user)
    elif request_path == "/api/imports":
        user, response = require_admin(environ)
        if response:
            status, headers, body = response
        else:
            with connect() as conn:
                rows = conn.execute(
                    """
                    SELECT ib.id, ib.outlet, ib.filename, ib.imported_count, ib.created_at, u.username
                    FROM import_batches ib
                    LEFT JOIN users u ON u.id = ib.user_id
                    ORDER BY ib.created_at DESC
                    LIMIT 50
                    """
                ).fetchall()
            status, headers, body = json_response({"imports": [dict(row) for row in rows]})
    elif request_path == "/api/import" and method == "POST":
        user, response = require_editor(environ)
        if response:
            status, headers, body = response
        else:
            query = parse_qs(environ.get("QUERY_STRING", ""))
            outlet = query.get("outlet", ["SHIMA"])[0].upper()
            if outlet not in OUTLETS:
                status, headers, body = json_response({"error": "Outlet no válido"}, "400 Bad Request")
            else:
                size = int(environ.get("CONTENT_LENGTH") or "0")
                if size <= 0:
                    status, headers, body = json_response({"error": "Excel vacío"}, "400 Bad Request")
                    start_response(status, headers)
                    return [body]
                if size > MAX_UPLOAD_BYTES:
                    status, headers, body = json_response({"error": "Excel demasiado grande"}, "413 Payload Too Large")
                    start_response(status, headers)
                    return [body]
                payload = environ["wsgi.input"].read(size)
                filename = safe_filename(clean(environ.get("HTTP_X_FILENAME", "")) or f"{outlet}.xlsx")
                if not filename.lower().endswith(".xlsx"):
                    status, headers, body = json_response({"error": "Solo se aceptan archivos .xlsx"}, "400 Bad Request")
                    start_response(status, headers)
                    return [body]
                outlet_dir = IMPORTS_DIR / outlet.replace(" ", "_")
                outlet_dir.mkdir(parents=True, exist_ok=True)
                timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
                target = outlet_dir / f"{timestamp}-{filename}"
                target.write_bytes(payload)
                records = parse_workbook(target, outlet)
                if not records:
                    status, headers, body = json_response(
                        {"error": "El Excel se ha leido, pero no se ha encontrado ningun vino importable. Revisa columnas y pestañas."},
                        "400 Bad Request",
                    )
                    start_response(status, headers)
                    return [body]
                count = import_records(records, outlet)
                with connect() as conn:
                    conn.execute(
                        """
                        INSERT INTO import_batches(outlet, filename, imported_count, user_id, created_at)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (outlet, filename, count, user["id"], now_iso()),
                    )
                status, headers, body = json_response({"imported": count, "outlet": outlet, "summary": database_summary()})
    else:
        user, response = require_user(environ)
        if response:
            status, headers, body = response
        else:
            file_path = static_file_for(request_path)
            if file_path is None:
                status, headers, body = "404 Not Found", [("Content-Type", "text/plain")], b"Not found"
            else:
                status, headers, body = static_response(file_path)

    start_response(status, headers)
    return [body]


if __name__ == "__main__":
    seed_from_excel()
    port = 8000
    print(f"WOS running on http://127.0.0.1:{port}")
    if not os.environ.get("WOS_ADMIN_PASSWORD") and ADMIN_PASSWORD_PATH.exists():
        print("Usuario admin: admin")
        print(f"Contraseña admin: {ADMIN_PASSWORD_PATH.read_text(encoding='utf-8').strip()}")
    make_server("0.0.0.0", port, app).serve_forever()
